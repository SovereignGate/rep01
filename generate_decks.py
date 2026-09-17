import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# -------------------------------------------------------------
# 1. BUILD PPTX PRESENTATION
# -------------------------------------------------------------
def create_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(15, 23, 42)
    AZURE = RGBColor(0, 120, 212)
    GRAY = RGBColor(100, 116, 139)
    LIGHT_BG = RGBColor(248, 250, 252)
    WHITE = RGBColor(255, 255, 255)

    slides_data = [
        {
            "tag": "MICROSOFT AGENTATHON 2026",
            "title": "SovereignGate",
            "subtitle": "Enterprise Privacy Proxy & Multi-Agent Governance on Azure AI Foundry",
            "bullets": [
                "Built for regulated Contract Lifecycle Management (CLM) & SMEs",
                "Stops IP & PII leakage into untrusted model training loops",
                "Enforces regional data residency (UK DPA 2018 & EU GDPR) dynamically",
                "Seamless user experience: zero workflow friction, zero token leakage"
            ]
        },
        {
            "tag": "THE PROBLEM",
            "title": "The Shadow AI Dilemma for SMEs",
            "subtitle": "Balancing Frontier AI Productivity with Strict Regulatory Compliance",
            "bullets": [
                "Uncontrolled Browser LLM Usage: Employees paste sensitive client NDAs, master service agreements, and fee structures into free/public models.",
                "Data Residency Violations: Cross-border transmission breaches UK GDPR and Data Protection Act rules when routed through untrusted jurisdictions.",
                "IP Poisoning: Proprietary trade terms and party identities become permanent parts of public model re-training corpuses.",
                "Legacy DLP Failure: Traditional network firewalls simply block access, halting AI productivity and driving users toward riskier workarounds."
            ]
        },
        {
            "tag": "THE SOLUTION",
            "title": "The 3-Agent Sovereign Architecture",
            "subtitle": "Decoupled Ingestion, Compliant Dispatch, and Ephemeral Reconstruction",
            "bullets": [
                "Agent 1 (Compliance Sentinel): Ingests raw contract text, redacts sensitive business PII/IP into deterministic tokens (e.g., {{ORG_1}}, {{MONEY_2}}), and classifies residency rules.",
                "Ephemeral Session Vault: Secures un-redacted values strictly in local volatile memory; raw identities never touch the external network.",
                "Agent 2 (Sovereign Dispatcher): Negotiates deployment models across Azure AI Foundry nodes, routing exclusively to compliant sovereign regions (e.g., UK South).",
                "Agent 3 (Reconstructor & Audit): Ingests the model inference output, seamlessly restores real values, and writes a tamper-evident audit log."
            ]
        },
        {
            "tag": "LIVE DEMO HIGHLIGHTS",
            "title": "Streamlit Prototype Workflow",
            "subtitle": "Side-by-Side Zero-Friction Governance",
            "bullets": [
                "Raw Input: Master Services Agreement with £250,000 fee and London legal jurisdiction.",
                "Sentinel Redaction: Real-time substitution into {{ORG_1}}, {{ORG_2}}, and {{MONEY_1}} with UK_ONLY tier policy tagging.",
                "Foundry Execution: gpt-5-mini in UK South processes reasoning entirely over synthetic tokens with zero access to proprietary identities.",
                "Final Delivery: Fully restored, professional summary delivered to the user with a real-time compliance audit log."
            ]
        },
        {
            "tag": "JUDGING CRITERIA",
            "title": "Why SovereignGate Wins",
            "subtitle": "Evaluation Against Microsoft Agentathon Standards",
            "bullets": [
                "Innovation: Shifts DLP from static blocking to an active, autonomous multi-agent proxy with real-time token re-inflation.",
                "Usability: Fully native to Azure AI Foundry SDK. End users write standard prompts without learning specialized compliance syntax.",
                "Impact: Immediate elimination of GDPR fines (up to 4% global turnover) and zero intellectual property leakage.",
                "Vendor Agility: Open architecture built on Python that extends naturally to browser extensions, M365 Copilot, and external CLMs."
            ]
        },
        {
            "tag": "LOOKING FORWARD",
            "title": "Roadmap & Scaling",
            "subtitle": "Production Hardening over Post-Agentathon Timeline",
            "bullets": [
                "Vectorized Policy Knowledge Base: Azure AI Search integration for live clause checking against evolving EU/UK statutory guidelines.",
                "Browser Extension & Proxy Mode: Zero-click enterprise gateway automatically proxying any corporate AI request.",
                "Cosmos DB Audit Ledger: Long-term encrypted audit persistence for ISO 27001 and SOC 2 type II compliance verification.",
                "Automated Red-Teaming: Continuous evaluation using Azure AI Foundry safety evaluators to guarantee zero token leakage."
            ]
        }
    ]

    for slide_data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Header Box
        header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.8))
        tf = header_box.text_frame
        tf.word_wrap = True

        # Tag
        p_tag = tf.paragraphs[0]
        p_tag.text = slide_data["tag"]
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.color.rgb = AZURE

        # Title
        p_title = tf.add_paragraph()
        p_title.text = slide_data["title"]
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

        # Subtitle
        p_sub = tf.add_paragraph()
        p_sub.text = slide_data["subtitle"]
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = GRAY

        # Content Box
        content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(11.3), Inches(4.0))
        c_tf = content_box.text_frame
        c_tf.word_wrap = True

        for i, bullet in enumerate(slide_data["bullets"]):
            p = c_tf.paragraphs[0] if i == 0 else c_tf.add_paragraph()
            p.text = "•  " + bullet
            p.font.size = Pt(16)
            p.font.color.rgb = NAVY
            p.space_after = Pt(14)

    pptx_path = "SovereignGate_Pitch.pptx"
    prs.save(pptx_path)
    print(f"[+] PPTX generated successfully: {pptx_path}")


# -------------------------------------------------------------
# 2. BUILD PDF PRESENTATION (Landscape Slides)
# -------------------------------------------------------------
def create_pdf():
    pdf_path = "SovereignGate_Pitch.pdf"
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=landscape(letter),
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    tag_style = ParagraphStyle(
        'TagStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        textColor=colors.HexColor('#0078D4'),
        spaceAfter=4
    )
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=4
    )
    sub_style = ParagraphStyle(
        'SubStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=14
    )
    bullet_style = ParagraphStyle(
        'BulletStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=16,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=8
    )

    slides_data = [
        ("MICROSOFT AGENTATHON 2026", "SovereignGate", "Enterprise Privacy Proxy & Multi-Agent Governance on Azure AI Foundry", [
            "Built for regulated Contract Lifecycle Management (CLM) & SMEs",
            "Stops IP & PII leakage into untrusted model training loops",
            "Enforces regional data residency (UK DPA 2018 & EU GDPR) dynamically",
            "Seamless user experience: zero workflow friction, zero token leakage"
        ]),
        ("THE PROBLEM", "The Shadow AI Dilemma for SMEs", "Balancing Frontier AI Productivity with Strict Regulatory Compliance", [
            "Uncontrolled Browser LLM Usage: Employees paste sensitive client NDAs, master service agreements, and fee structures into free/public models.",
            "Data Residency Violations: Cross-border transmission breaches UK GDPR and Data Protection Act rules when routed through untrusted jurisdictions.",
            "IP Poisoning: Proprietary trade terms and party identities become permanent parts of public model re-training corpuses.",
            "Legacy DLP Failure: Traditional network firewalls simply block access, halting AI productivity and driving users toward riskier workarounds."
        ]),
        ("THE SOLUTION", "The 3-Agent Sovereign Architecture", "Decoupled Ingestion, Compliant Dispatch, and Ephemeral Reconstruction", [
            "Agent 1 (Compliance Sentinel): Ingests raw contract text, redacts sensitive business PII/IP into deterministic tokens (e.g., {{ORG_1}}, {{MONEY_2}}), and classifies residency rules.",
            "Ephemeral Session Vault: Secures un-redacted values strictly in local volatile memory; raw identities never touch the external network.",
            "Agent 2 (Sovereign Dispatcher): Negotiates deployment models across Azure AI Foundry nodes, routing exclusively to compliant sovereign regions (e.g., UK South).",
            "Agent 3 (Reconstructor & Audit): Ingests the model inference output, seamlessly restores real values, and writes a tamper-evident audit log."
        ]),
        ("LIVE DEMO HIGHLIGHTS", "Streamlit Prototype Workflow", "Side-by-Side Zero-Friction Governance", [
            "Raw Input: Master Services Agreement with £250,000 fee and London legal jurisdiction.",
            "Sentinel Redaction: Real-time substitution into {{ORG_1}}, {{ORG_2}}, and {{MONEY_1}} with UK_ONLY tier policy tagging.",
            "Foundry Execution: gpt-5-mini in UK South processes reasoning entirely over synthetic tokens with zero access to proprietary identities.",
            "Final Delivery: Fully restored, professional summary delivered to the user with a real-time compliance audit log."
        ]),
        ("JUDGING CRITERIA", "Why SovereignGate Wins", "Evaluation Against Microsoft Agentathon Standards", [
            "Innovation: Shifts DLP from static blocking to an active, autonomous multi-agent proxy with real-time token re-inflation.",
            "Usability: Fully native to Azure AI Foundry SDK. End users write standard prompts without learning specialized compliance syntax.",
            "Impact: Immediate elimination of GDPR fines (up to 4% global turnover) and zero intellectual property leakage.",
            "Vendor Agility: Open architecture built on Python that extends naturally to browser extensions, M365 Copilot, and external CLMs."
        ]),
        ("LOOKING FORWARD", "Roadmap & Scaling", "Production Hardening over Post-Agentathon Timeline", [
            "Vectorized Policy Knowledge Base: Azure AI Search integration for live clause checking against evolving EU/UK statutory guidelines.",
            "Browser Extension & Proxy Mode: Zero-click enterprise gateway automatically proxying any corporate AI request.",
            "Cosmos DB Audit Ledger: Long-term encrypted audit persistence for ISO 27001 and SOC 2 type II compliance verification.",
            "Automated Red-Teaming: Continuous evaluation using Azure AI Foundry safety evaluators to guarantee zero token leakage."
        ])
    ]

    elements = []

    for i, (tag, title, subtitle, bullets) in enumerate(slides_data):
        card_content = [
            [Paragraph(tag, tag_style)],
            [Paragraph(title, title_style)],
            [Paragraph(subtitle, sub_style)],
            [Spacer(1, 10)]
        ]
        for b in bullets:
            card_content.append([Paragraph(f"•  {b}", bullet_style)])

        table = Table(card_content, colWidths=[710])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#E2E8F0')),
            ('ROUNDEDCORNERS', [8, 8, 8, 8]),
            ('TOPPADDING', (0, 0), (-1, -1), 14),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 14),
            ('LEFTPADDING', (0, 0), (-1, -1), 24),
            ('RIGHTPADDING', (0, 0), (-1, -1), 24),
        ]))

        elements.append(table)
        if i < len(slides_data) - 1:
            from reportlab.platypus import PageBreak
            elements.append(PageBreak())

    doc.build(elements)
    print(f"[+] PDF generated successfully: {pdf_path}")


if __name__ == "__main__":
    create_pptx()
    create_pdf()
